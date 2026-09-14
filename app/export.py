"""
Export layer for extracted invoice data -- lets non-technical users
(accounting/finance teams without developers) consume the API's output
without touching JSON. Builds an .xlsx workbook and a .csv, both from
the same dict returned by app.invoice.extract_invoice_fields.
"""
import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

SUMMARY_FIELD_LABELS = [
    ("vendor", "Vendor / Seller Name"),
    ("seller_address", "Seller Address"),
    ("seller_phone", "Seller Phone"),
    ("seller_gstin", "Seller GSTIN"),
    ("buyer_name", "Buyer Name"),
    ("buyer_address", "Buyer Address"),
    ("buyer_phone", "Buyer Phone"),
    ("buyer_gstin", "Buyer GSTIN"),
    ("invoice_number", "Invoice Number"),
    ("invoice_type", "Invoice Type"),
    ("financial_year", "Financial Year"),
    ("date", "Invoice Date"),
    ("taxable_value", "Taxable Value"),
    ("discount_amount", "Discount"),
    ("total_amount", "Total Amount"),
    ("amount_in_words", "Amount in Words"),
    ("tax_amount", "Total Tax"),
    ("cgst_amount", "CGST"),
    ("sgst_amount", "SGST"),
    ("igst_amount", "IGST"),
    ("vat_amount", "VAT"),
    ("sales_tax_amount", "Sales Tax"),
    ("gst_amount", "GST"),
    ("hst_amount", "HST"),
    ("pst_amount", "PST"),
    ("qst_amount", "QST"),
    ("consumption_tax_amount", "Consumption Tax"),
    ("tax_split_source", "Tax Split Source"),
    ("confidence", "Extraction Confidence"),
    ("page_count", "Page Count"),
]

LINE_ITEM_COLUMNS = [
    ("description", "Description"),
    ("hsn_sac", "HSN/SAC"),
    ("quantity", "Quantity"),
    ("unit_price", "Unit Price"),
    ("gst_rate", "GST Rate (%)"),
    ("amount", "Amount"),
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LABEL_FONT = Font(bold=True)


def build_invoice_xlsx(data: dict) -> bytes:
    wb = Workbook()

    summary = wb.active
    summary.title = "Invoice Summary"
    summary["A1"] = "Field"
    summary["B1"] = "Value"
    for cell in (summary["A1"], summary["B1"]):
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    row = 2
    for key, label in SUMMARY_FIELD_LABELS:
        val = data.get(key)
        summary.cell(row=row, column=1, value=label).font = LABEL_FONT
        summary.cell(row=row, column=2, value=val if val is not None else "")
        row += 1
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 55
    for r in range(1, row):
        summary.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    items_sheet = wb.create_sheet("Line Items")
    for col_idx, (_, label) in enumerate(LINE_ITEM_COLUMNS, start=1):
        cell = items_sheet.cell(row=1, column=col_idx, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    line_items = data.get("line_items") or []
    for r, item in enumerate(line_items, start=2):
        for c, (key, _) in enumerate(LINE_ITEM_COLUMNS, start=1):
            items_sheet.cell(row=r, column=c, value=item.get(key, ""))
    for col_idx in range(1, len(LINE_ITEM_COLUMNS) + 1):
        items_sheet.column_dimensions[get_column_letter(col_idx)].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_invoice_csv(data: dict) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)

    writer.writerow(["Field", "Value"])
    for key, label in SUMMARY_FIELD_LABELS:
        val = data.get(key)
        writer.writerow([label, val if val is not None else ""])

    writer.writerow([])
    writer.writerow([label for _, label in LINE_ITEM_COLUMNS])
    for item in (data.get("line_items") or []):
        writer.writerow([item.get(key, "") for key, _ in LINE_ITEM_COLUMNS])

    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 CSV correctly
